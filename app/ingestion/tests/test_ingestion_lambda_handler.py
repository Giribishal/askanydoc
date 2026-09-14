"""Unit tests for the first S3 ingestion Lambda milestone."""

import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

# The runtime directory is deployed to Lambda; add it here so the local test can
# import the handler without turning the learning and test folders into runtime code.
RUNTIME_DIR = Path(__file__).resolve().parents[1] / "runtime"
SHARED_DIR = Path(__file__).resolve().parents[2] / "shared"
sys.path.insert(0, str(SHARED_DIR))
sys.path.insert(0, str(RUNTIME_DIR))

from extract_text import extract_docx_from_bytes, extract_utf8_text_from_bytes
from ingestion_lambda_handler import detect_source_type, handler, route_record


def s3_event(key: str, *, size: int | None = None, version_id: str | None = None) -> dict:
    object_details = {"key": key}
    if size is not None:
        object_details["size"] = size
    if version_id is not None:
        object_details["versionId"] = version_id
    return {
        "Records": [{
            "eventSource": "aws:s3",
            "s3": {
                "bucket": {"name": "askanydoc-documents-prod-apse2"},
                "object": object_details,
            },
        }]
    }


class IngestionRoutingTests(unittest.TestCase):
    def test_supported_extensions_are_case_insensitive(self) -> None:
        expected = {"a.PDF": "pdf", "b.docx": "docx", "c.TXT": "txt", "d.md": "md"}
        for key, source_type in expected.items():
            with self.subTest(key=key):
                self.assertEqual(detect_source_type(key), source_type)

    def test_url_encoded_key_is_decoded_and_routed(self) -> None:
        result = route_record(s3_event("uploads/HR+Policy.docx")["Records"][0])
        self.assertEqual(result["status"], "routed")
        self.assertEqual(result["source_name"], "HR Policy.docx")
        self.assertEqual(result["source_type"], "docx")
        self.assertEqual(result["extractor"], "extract_docx")
        self.assertEqual(result["source_uri"], "s3://askanydoc-documents-prod-apse2/uploads/HR Policy.docx")
        self.assertTrue(result["document_id"])

    def test_unsupported_file_is_ignored_cleanly(self) -> None:
        result = handler(s3_event("uploads/report.xlsx"), None)["results"][0]
        self.assertEqual(result["status"], "ignored")
        self.assertEqual(result["reason"], "unsupported_file_type")

    def test_oversized_document_is_rejected_before_download(self) -> None:
        result = handler(
            s3_event("uploads/large.pdf", size=10 * 1024 * 1024 + 1), None
        )["results"][0]
        self.assertEqual(result["status"], "ignored")
        self.assertEqual(result["reason"], "document_too_large")

    def test_pdf_is_downloaded_and_extracted_from_s3_bytes(self) -> None:
        # Build a valid one-page PDF in memory so the test needs no fixture or AWS access.
        pdf_output = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.write(pdf_output)

        sample_chunk = {
            "document_id": "document-id",
            "source_name": "example.pdf",
            "source_uri": "s3://askanydoc-documents-prod-apse2/uploads/example.pdf",
            "source_sha256": "source-hash",
            "text": "sample text",
            "location": {"page_number": 1},
            "chunk_id": "chunk-id",
            "chunk_index": 0,
        }
        embedded_chunk = {**sample_chunk, "embedding": [0.0] * 1024}

        with (
            patch("ingestion_lambda_handler.s3_client.get_object") as get_object,
            patch("ingestion_lambda_handler.chunk_pages", return_value=[sample_chunk]),
            patch("ingestion_lambda_handler.embed_chunks", return_value=[embedded_chunk]),
            patch("ingestion_lambda_handler.replace_document_chunks") as replace_chunks,
        ):
            get_object.return_value = {"Body": BytesIO(pdf_output.getvalue())}
            result = handler(
                s3_event("uploads/example.pdf", version_id="document-version-1"), None
            )["results"][0]

        get_object.assert_called_once_with(
            Bucket="askanydoc-documents-prod-apse2",
            Key="uploads/example.pdf",
            VersionId="document-version-1",
        )
        self.assertEqual(result["status"], "extracted")
        self.assertEqual(result["page_count"], 1)
        self.assertEqual(result["non_empty_page_count"], 0)
        self.assertEqual(result["character_count"], 0)
        self.assertEqual(result["chunk_count"], 1)
        self.assertTrue(result["source_sha256"])
        replace_chunks.assert_called_once()

    def test_utf8_text_and_markdown_are_extracted(self) -> None:
        for source_name in ("notes.txt", "guide.md"):
            with self.subTest(source_name=source_name):
                records = extract_utf8_text_from_bytes(
                    b"Heading\nUseful content", source_name, f"s3://bucket/{source_name}"
                )
                self.assertEqual(records[0]["text"], "Heading\nUseful content")
                self.assertEqual(records[0]["status"], "extracted")

    def test_docx_paragraphs_and_tables_are_extracted(self) -> None:
        output = BytesIO()
        document = Document()
        document.add_paragraph("Policy heading")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Owner"
        table.cell(0, 1).text = "Security"
        document.save(output)

        records = extract_docx_from_bytes(
            output.getvalue(), "policy.docx", "s3://bucket/policy.docx"
        )
        self.assertIn("Policy heading", records[0]["text"])
        self.assertIn("Owner | Security", records[0]["text"])


if __name__ == "__main__":
    unittest.main()
